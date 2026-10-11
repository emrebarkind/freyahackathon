"""Case materials: folder and zip input, Word/PDF/Excel read with their file, page and sheet kept, unreadable files
reported (never skipped silently), and the bank mock URL applied to every bank call of the compiled workflow."""
import json, pathlib, zipfile
import docx, openpyxl
from builder import materials
from builder import compile as comp
from builder import build

ROOT = pathlib.Path(__file__).resolve().parent.parent
OA = ROOT / "kit" / "build-the-builder" / "tools" / "openapi.json"


def make_case(d):
    (d / "BRIEF.md").write_text("# Brief\nKısmi ödeme kabul edilir. Ödeme günü 30 Ekim'e kadar olabilir.", encoding="utf-8")
    doc = docx.Document(); doc.add_heading("Servis sözleşmesi", 1); doc.add_paragraph("create_promise tutarı kuruş olarak alır.")
    t = doc.add_table(rows=1, cols=2); t.rows[0].cells[0].text = "kod"; t.rows[0].cells[1].text = "REF"
    doc.save(d / "sozlesme.docx")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Müşteriler"; ws.append(["id", "borç"]); ws.append(["C1", 1250.5]); wb.save(d / "test_musterileri.xlsx")
    (d / "akis.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 20)
    (d / "openapi.json").write_text(OA.read_text())


def test_folder_reads_every_format_with_its_origin(tmp_path):
    make_case(tmp_path)
    api, api_path, docs, mock, unread, notes = materials.gather([str(tmp_path), "https://mock.example.workers.dev/"])
    names = [n for n, _ in docs]
    assert api_path.endswith("openapi.json") and mock == "https://mock.example.workers.dev"
    assert {"BRIEF.md", "sozlesme.docx", "test_musterileri.xlsx"} <= set(names)
    text = {n: t for n, t in docs}
    assert "kuruş" in text["sozlesme.docx"] and "REF" in text["sozlesme.docx"]          # table kept
    assert "[sheet Müşteriler]" in text["test_musterileri.xlsx"] and "1250.5" in text["test_musterileri.xlsx"]
    assert "30 Ekim" in text["BRIEF.md"]                                                 # Turkish kept
    assert any(u.startswith("akis.png") for u in unread)                                # not skipped silently


def test_zip_input_is_extracted_and_the_source_is_unchanged(tmp_path):
    case = tmp_path / "case"; case.mkdir(); make_case(case)
    z = tmp_path / "case.zip"
    with zipfile.ZipFile(z, "w") as f:
        for p in case.iterdir():
            f.write(p, "case/" + p.name)
    before = z.read_bytes()
    api, api_path, docs, mock, unread, notes = materials.gather([str(z)])
    assert api is not None and any(n.endswith("BRIEF.md") for n, _ in docs) and z.read_bytes() == before


def test_pdf_pages_are_labelled(tmp_path):
    import subprocess
    (tmp_path / "a.md").write_text("# Duyuru\n\nBirinci sayfa.\n\n\\newpage\n\nİkinci sayfa.", encoding="utf-8")
    r = subprocess.run(["pandoc", str(tmp_path / "a.md"), "-o", str(tmp_path / "a.pdf")], capture_output=True)
    if r.returncode != 0:
        import pytest; pytest.skip("pandoc cannot make a PDF here")
    text, problem = materials.read(str(tmp_path / "a.pdf"))
    assert problem is None and "[page 1]" in text


def test_the_mock_url_reaches_every_bank_call_including_chained_copies():
    spec = json.load(open(ROOT / "eval" / "clean-runs" / "cand8-a" / "cand8-a.spec.json"))
    oa = json.load(open(OA)); oa["servers"] = [{"url": "https://mock.example.workers.dev"}]
    s = comp.apply_defaults(spec, oa); wf = comp.compile_workflow(s, oa)
    urls = [a["function"]["parameters"]["properties"]["url"]["default"] for n in wf["nodes"].values()
            for a in n.get("pre_actions") or [] if "url" in ((a.get("function") or {}).get("parameters") or {}).get("properties", {})]
    assert len(urls) >= 10 and all(u.startswith("https://mock.example.workers.dev/") for u in urls)
    assert any(n.endswith("_2") for n in wf["nodes"])                                    # chained copies exist and are covered
    build.COMPILED[0], build.BANK_BASE[0] = wf, "https://mock.example.workers.dev"
    assert build.bank_url_gaps(s) == []
    build.BANK_BASE[0] = "https://other.example"
    assert build.bank_url_gaps(s)
    build.COMPILED[0] = build.BANK_BASE[0] = None
